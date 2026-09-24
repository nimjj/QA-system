#!/usr/bin/env python3
"""
export_results_to_excel.py

Reads all QA test result JSON files from 'inputs/Test (results)'
and generates / updates an Excel spreadsheet in the project root directory.

Columns:
  1. File Name: The name of each result JSON file
  2. v 8.1 score: The evaluated final score
  3. Complete Body: The full JSON content of the file
"""

import os
import sys
import json
import argparse
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: 'openpyxl' is required to generate Excel files.")
    print("Please install it using: pip install openpyxl")
    sys.exit(1)


# Default configurations
DEFAULT_INPUT_DIR = os.path.join("inputs", "Test (results)")
DEFAULT_OUTPUT_FILE = "test_results.xlsx"

# Header labels
COL_FILE_NAME = "File Name"
COL_SCORE = "v 8.1 score"
COL_BODY = "Complete Body"


def parse_result_file(file_path: Path):
    """
    Parses a result JSON file and extracts:
      - filename
      - final score
      - raw formatted JSON content
    """
    with open(file_path, "r", encoding="utf-8") as f:
        raw_content = f.read()

    try:
        data = json.loads(raw_content)
    except json.JSONDecodeError:
        data = {}

    # Extract final score safely
    final_score = None
    if isinstance(data, dict):
        result_obj = data.get("result")
        if isinstance(result_obj, dict):
            final_score = result_obj.get("final_score")
        elif "final_score" in data:
            final_score = data.get("final_score")

    # Format the complete body nicely if valid JSON, otherwise keep raw content
    try:
        if data:
            body_content = json.dumps(data, indent=2, ensure_ascii=False)
        else:
            body_content = raw_content
    except Exception:
        body_content = raw_content

    return {
        "file_name": file_path.name,
        "score": final_score,
        "body": body_content,
        "status": data.get("status", "unknown") if isinstance(data, dict) else "unknown"
    }


def find_result_files(input_dir: Path, recursive: bool = False):
    """
    Finds all result JSON files inside the given directory.
    """
    if not input_dir.exists():
        print(f"Directory not found: {input_dir}")
        return []

    pattern = "**/*.json" if recursive else "*.json"
    files = sorted(
        [f for f in input_dir.glob(pattern) if f.is_file() and not f.name.startswith(".")],
        key=lambda p: p.name
    )
    return files


def create_or_update_excel(records: list, output_path: Path, col_order: str = "score_first"):
    """
    Writes records to an Excel workbook with styling and auto-formatting.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Results"

    # Configure column layout based on preference
    if col_order == "body_first":
        headers = [COL_FILE_NAME, COL_BODY, COL_SCORE]
    else:  # score_first (recommended for usability)
        headers = [COL_FILE_NAME, COL_SCORE, COL_BODY]

    # Style definitions
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

    data_font = Font(name="Segoe UI", size=10)
    score_font = Font(name="Segoe UI", size=11, bold=True)
    body_font = Font(name="Consolas", size=9)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Alternate row shading
    even_row_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    odd_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Write Header Row
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        if header == COL_FILE_NAME or header == COL_BODY:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write Data Rows
    for row_idx, rec in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20
        row_fill = even_row_fill if row_idx % 2 == 0 else odd_row_fill

        if col_order == "body_first":
            row_data = [rec["file_name"], rec["body"], rec["score"]]
        else:
            row_data = [rec["file_name"], rec["score"], rec["body"]]

        for col_idx, (val, header) in enumerate(zip(row_data, headers), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.fill = row_fill

            if header == COL_FILE_NAME:
                cell.font = data_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif header == COL_SCORE:
                cell.font = score_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(val, (int, float)):
                    cell.number_format = "0.0"
            elif header == COL_BODY:
                cell.font = body_font
                # wrap_text is False so rows don't explode to 1000px height in Excel;
                # the full formatted multiline JSON is preserved in the cell
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)

    # Set Column Widths
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        if header == COL_FILE_NAME:
            ws.column_dimensions[col_letter].width = 54
        elif header == COL_SCORE:
            ws.column_dimensions[col_letter].width = 16
        elif header == COL_BODY:
            ws.column_dimensions[col_letter].width = 65

    # Freeze header row
    ws.freeze_panes = "A2"

    # Enable Excel Auto-Filter on header row
    ws.auto_filter.ref = ws.dimensions

    # Enable grid lines visibility
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    # Save to disk with lock handling
    try:
        wb.save(output_path)
        return True, None
    except PermissionError:
        backup_path = output_path.with_name(f"{output_path.stem}_new{output_path.suffix}")
        try:
            wb.save(backup_path)
            return False, backup_path
        except Exception as e:
            return False, str(e)


def main():
    parser = argparse.ArgumentParser(
        description="Extract QA test result JSON files and update an Excel spreadsheet."
    )
    parser.add_argument(
        "--input-dir", "-i",
        default=DEFAULT_INPUT_DIR,
        help=f"Directory containing test result JSON files (default: '{DEFAULT_INPUT_DIR}')"
    )
    parser.add_argument(
        "--output", "-o",
        default=DEFAULT_OUTPUT_FILE,
        help=f"Output Excel file path (default: '{DEFAULT_OUTPUT_FILE}')"
    )
    parser.add_argument(
        "--recursive", "-r",
        action="store_true",
        help="Search for JSON files recursively in subdirectories"
    )
    parser.add_argument(
        "--col-order",
        choices=["score_first", "body_first"],
        default="score_first",
        help="Column layout order: 'score_first' (File Name, v 8.1 score, Complete Body) or 'body_first' (File Name, Complete Body, v 8.1 score)"
    )

    args = parser.parse_args()

    # Determine paths relative to QA-system project root
    base_dir = Path(__file__).resolve().parent
    input_path = Path(args.input_dir)
    if not input_path.is_absolute():
        input_path = base_dir / input_path

    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = base_dir / output_path

    print(f"Scanning for JSON result files in: {input_path}")
    files = find_result_files(input_path, recursive=args.recursive)

    if not files:
        print(f"No JSON files found in '{input_path}'.")
        return

    print(f"Found {len(files)} result JSON files. Processing data...")

    records = []
    scores = []
    for f in files:
        rec = parse_result_file(f)
        records.append(rec)
        if rec["score"] is not None and isinstance(rec["score"], (int, float)):
            scores.append(rec["score"])

    # Update Excel
    success, backup_or_err = create_or_update_excel(records, output_path, col_order=args.col_order)

    if success:
        print(f"\n[SUCCESS] Successfully updated Excel file:")
        print(f"  -> Path: {output_path.resolve()}")
    else:
        if isinstance(backup_or_err, Path):
            print(f"\n[WARNING] Could not overwrite '{output_path.name}' (file is open in Excel).")
            print(f"  -> Saved changes to backup file: {backup_or_err.resolve()}")
            print(f"  -> Please close '{output_path.name}' in Excel to overwrite it on subsequent runs.")
        else:
            print(f"\n[ERROR] Failed to save Excel file: {backup_or_err}")
            return

    # Print summary statistics
    print("\n--- Summary ---")
    print(f"  Total tests processed : {len(records)}")
    if scores:
        avg_score = sum(scores) / len(scores)
        perfect_count = sum(1 for s in scores if s >= 100.0)
        fail_count = sum(1 for s in scores if s == 0.0)
        print(f"  Average v 8.1 score   : {avg_score:.2f}")
        print(f"  Perfect scores (100)  : {perfect_count}")
        print(f"  Auto-fail scores (0)  : {fail_count}")
    print("----------------\n")


if __name__ == "__main__":
    main()
