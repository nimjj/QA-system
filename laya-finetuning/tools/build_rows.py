"""Turn hand-written call fragments into training rows, using the SAME input builder as production.

Fragments live in data/fragments/claude/*.txt in a compact format:

    @@ id=para_001 check=paraphrasing label=PASS pair=para_p01 industry=telecom persona=warm hard=restate_vs_echo name=Maya
    C(-12): My internet keeps dropping every evening around eight.
    A*(+6): So it drops each evening around eight, is that right?      <- * marks the decision line(s)
    # comment lines and blank lines are ignored;  "@gap 45" adds 45 s of silence before the next turn

Speakers are A (Agent) and C (Customer); (n) is the per-turn sentiment score (defaults: agent +6, customer -4).

For every fragment this script: applies speech-to-text roughening to ~30% of PAIRS (never to one member of a pair),
runs qa-analysis's build_laya_inputs(), takes the snippet for the fragment's check, and DROPS the fragment (logged
with a reason) if production would not send it to Laya as intended: the check's gate is not reached, there is more
than one snippet, the decision line is not in it, or it does not fit the model window with the real tokenizer.
Rows carry `inputs_sha256`; the build refuses to run if qa-analysis no longer matches data/frozen/checks_frozen.json.

Usage (qa-analysis venv):
    python tools/build_rows.py --fragments "data/fragments/claude/pilot_*.txt" --out data/rows/pilot.jsonl
    python tools/build_rows.py ... --sheet data/review/pilot_review.xlsx        # also the human skim sheet
"""
import argparse
import collections
import glob
import hashlib
import json
import os
import random
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import paths  # noqa: E402
from roughen import roughen_text  # noqa: E402

ROUGH_PERCENT = 30
LINE_RE = re.compile(r"^(A|C)(\*)?(?:\(([+-]?\d+(?:\.\d+)?)\))?:\s*(.*)$")


# ------------------------------------------------------------------ parsing
def parse_turns(lines, src):
    turns, pending_gap = [], 0
    for n, line in lines:
        if line.startswith("@gap"):
            pending_gap = int(line.split()[1])
            continue
        m = LINE_RE.match(line)
        if not m:
            raise ValueError(f"{src}:{n}: cannot read line: {line[:60]}")
        spk, focus, sent, text = m.groups()
        turns.append({"speaker": "Agent" if spk == "A" else "Customer", "text": text.strip(), "focus": bool(focus),
                      "sentiment": float(sent) if sent is not None else (6.0 if spk == "A" else -4.0), "gap": pending_gap})
        pending_gap = 0
    return turns


def parse_attrs(header, src):
    attrs = {k: v.strip('"') for k, v in re.findall(r"(\w+)=((?:\"[^\"]*\")|\S+)", header)}
    for need in ("id", "check"):
        if need not in attrs:
            raise ValueError(f"{src}: fragment header lacks {need}=")
    return attrs


def parse_file(path):
    """Two block kinds.
    Single:  "@@ id=.. check=.. label=PASS|FAIL [pair=..] ..." followed by turn lines.
    Pair:    "@@pair id=.. check=.. industry=.. persona=.. hard=.." with shared turn lines first, then
             "--- PASS" and "--- FAIL" sections (the part that differs), then optional "--- END" shared closing lines.
             Expands to two fragments, <id>_p (PASS) and <id>_f (FAIL), both with pair=<id> and the same persona."""
    blocks, cur = [], None
    for n, raw in enumerate(open(path, encoding="utf-8"), 1):
        line = raw.rstrip("\n")
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        if line.startswith("@@"):
            cur = {"header": line, "n": n, "lines": []}
            blocks.append(cur)
        elif cur is None:
            raise ValueError(f"{path}:{n}: line before any @@ header")
        else:
            cur["lines"].append((n, line))
    frags = []
    for blk in blocks:
        src = f"{os.path.basename(path)}:{blk['n']}"
        if blk["header"].startswith("@@pair"):
            attrs = parse_attrs(blk["header"][6:], src)
            sections, name = {"": []}, ""
            for n, line in blk["lines"]:
                m = re.match(r"^---\s*(PASS|FAIL|END)\s*$", line)
                if m:
                    name = m.group(1)
                    sections.setdefault(name, [])
                else:
                    sections[name].append((n, line))
            if "PASS" not in sections or "FAIL" not in sections:
                raise ValueError(f"{src}: a pair needs both --- PASS and --- FAIL sections")
            for label, suffix in (("PASS", "_p"), ("FAIL", "_f")):
                frags.append({**attrs, "id": attrs["id"] + suffix, "pair": attrs["id"], "label": label, "src": src,
                              "turns": parse_turns(sections[""] + sections[label] + sections.get("END", []), src)})
        else:
            attrs = parse_attrs(blk["header"][2:], src)
            if "label" not in attrs:
                raise ValueError(f"{src}: fragment header lacks label=")
            frags.append({**attrs, "src": src, "turns": parse_turns(blk["lines"], src)})
    return frags


def use_rough(frag):
    if frag.get("rough") in ("0", "1"):
        return frag["rough"] == "1"
    key = frag.get("pair") or frag["id"]          # per pair, so roughness can never leak the label
    return int(hashlib.md5(key.encode()).hexdigest()[:8], 16) % 100 < ROUGH_PERCENT


def to_transcript(frag):
    rough = use_rough(frag)
    rng = random.Random(int(hashlib.md5(frag["id"].encode()).hexdigest()[:8], 16))
    t, out, focus_idx = 0, [], []
    for i, turn in enumerate(frag["turns"]):
        text = roughen_text(turn["text"], rng) if rough else turn["text"]
        t += turn["gap"]
        out.append({"speaker": turn["speaker"], "text": text, "start_time_sec": t, "sentiment_score": turn["sentiment"]})
        if turn["focus"]:
            focus_idx.append(i)
        t += max(6, int(len(text.split()) * 0.45)) + 2
    return out, focus_idx, rough


# ------------------------------------------------------------------ snippet selection per check
def select(check, inp, transcript, focus_idx, qa_rule_engine):
    """(snippet, None) or (None, reason): what production sends Laya for this fragment's check."""
    focus_texts = [transcript[i]["text"] for i in focus_idx]
    if check == "rapport":
        if not inp.rapport:
            return None, "no agent lines"
        chunks = inp.rapport
        if focus_texts:
            hit = [c for c in chunks if any(f"Agent: {f}" in c for f in focus_texts)]
            if len(hit) != 1:
                return None, "decision line not in exactly one chunk"
            return hit[0], None
        if len(chunks) != 1:
            return None, f"{len(chunks)} chunks and no decision line marked"
        return chunks[0], None
    if check == "paraphrasing":
        if inp.paraphrasing.split("AGENT:", 1)[-1].strip() == "" or inp.paraphrasing.startswith("CUSTOMER: \n"):
            return None, "empty customer or agent part"
        return inp.paraphrasing, None
    if check == "active_listening":
        if not inp.active_listening:
            return None, "gate not reached: no question pair above similarity 0.55 and no 'like I said' reply"
        blocks = inp.active_listening.split("\n\n")
        if len(blocks) != 1:
            return None, f"{len(blocks)} pairs found, expected exactly 1"
        if focus_idx:
            qs = qa_rule_engine.question_sentences(transcript[focus_idx[-1]]["text"])
            if not any(q in blocks[0] for q in qs):
                return None, "decision question is not the second question of the pair"
        return blocks[0], None
    if check == "resolved":
        if focus_texts and not all(f in inp.resolved for f in focus_texts):
            return None, "decision line is outside the last 10 lines"
        return inp.resolved, None
    if check == "empathy":
        if len(inp.empathy) != 1:
            return None, f"{len(inp.empathy)} empathy snippets, expected exactly 1 (gate not reached or extra upset turns)"
        if focus_texts and not inp.empathy[0].startswith("Customer: " + focus_texts[0]):
            return None, "snippet does not start at the intended upset turn"
        return inp.empathy[0], None
    if check == "ownership":
        return (inp.ownership, None) if inp.ownership else (None, "gate not reached: no deflection phrase in an agent line")
    if check == "probing":
        return (inp.probing, None) if inp.probing is not None else (None, "agent asked no questions (production fails it without the model)")
    if check in ("hostility", "escalation"):
        return inp.batch_full, None
    return None, f"unknown check {check}"


# ------------------------------------------------------------------ tokenizer / window
def load_tokenizer():
    sys.path.insert(0, paths.QA_ROOT)
    snap = paths.base_snapshot()
    if snap is None or not os.path.isdir(snap):
        sys.exit("STOP: base Laya weights not found. Set LAYA_HF_HOME to the Hugging Face cache that holds "
                 "convaiinnovations/laya revision " + paths.BASE_REVISION + " (see docs/NEW_PC_SETUP.md).")
    cfg = json.load(open(os.path.join(snap, "rl_agent_config.json"), encoding="utf-8"))
    from laya.agent import _load_tokenizer
    return _load_tokenizer(os.path.join(snap, "tokenizer"), cfg), cfg


class Builder:
    """Everything needed to turn one fragment into one row: the frozen questions, the real tokenizer, and
    qa-analysis's input builder. Used by this script and by tools/import_gemini.py."""

    def __init__(self, source="claude"):
        self.source = source
        sys.path.insert(0, paths.QA_ROOT)
        chk = subprocess.run([sys.executable, os.path.join(HERE, "export_checks.py"), "--check"], capture_output=True, text=True)
        print(chk.stdout.strip())
        if chk.returncode != 0:
            sys.exit("STOP: qa-analysis no longer matches data/frozen/checks_frozen.json. Regenerate the data after a deliberate re-freeze.")
        self.frozen = json.load(open(os.path.join(REPO, "data", "frozen", "checks_frozen.json"), encoding="utf-8"))
        self.checks = self.frozen["frozen"]["checks"]
        from src.services.laya_inputs import build_laya_inputs
        import src.services.rule_engine as rule_engine
        import laya.common as lc
        import laya.train as lt
        self.build_laya_inputs, self.rule_engine, self.lc, self.lt = build_laya_inputs, rule_engine, lc, lt
        self.tok, cfg = load_tokenizer()
        self.max_len, self.head_max_len = cfg["max_len"], cfg["head_max_len"]

    def row(self, f):
        """(row, None) or (None, {"id", "reason", ...}) for one parsed fragment."""
        check = f["check"]
        if check not in self.checks:
            return None, {"id": f["id"], "reason": f"unknown check {check}", "src": f.get("src")}
        transcript, focus_idx, rough = to_transcript(f)
        inp = self.build_laya_inputs(transcript, f.get("name"))
        snippet, why = select(check, inp, transcript, focus_idx, self.rule_engine)
        if snippet is None:
            return None, {"id": f["id"], "check": check, "reason": why, "src": f.get("src")}
        q = self.checks[check]["question"]
        room = self.lc.state_room(self.tok, self.lt.to_internal(check, q), self.max_len, self.head_max_len)
        n_tok = len(self.lt.encode_state(self.tok, snippet, 10 ** 6))
        if n_tok > room:
            return None, {"id": f["id"], "check": check, "reason": f"{n_tok} tokens > window room {room}", "src": f.get("src")}
        label = f["label"].upper()
        return {
            "state": snippet,
            "questions": {check: q},
            "expected": {check: label == self.checks[check]["yes_means"]},
            "meta": {"id": f["id"], "call_id": f["id"], "pair_id": f.get("pair", f["id"]), "check": check, "label": label,
                     "industry": f.get("industry"), "persona": f.get("persona"), "hard": f.get("hard"),
                     "rough": rough, "source": self.source, "tokens": n_tok, "room": room,
                     "inputs_sha256": self.frozen["inputs_sha256"]},
        }, None


def report(frags, rows, dropped):
    print(f"\nfragments {len(frags)} -> rows {len(rows)}, dropped {len(dropped)}")
    by = collections.defaultdict(collections.Counter)
    for r in rows:
        by[r["meta"]["check"]][r["meta"]["label"]] += 1
    print(f"{'check':18s} {'PASS':>5s} {'FAIL':>5s}")
    for c in sorted(by):
        print(f"{c:18s} {by[c]['PASS']:5d} {by[c]['FAIL']:5d}")
    print(f"roughened: {sum(r['meta']['rough'] for r in rows)}/{len(rows)}   max tokens {max((r['meta']['tokens'] for r in rows), default=0)}")
    ln = collections.defaultdict(lambda: collections.defaultdict(list))
    for r in rows:
        ln[r["meta"]["check"]][r["meta"]["label"]].append(r["meta"]["tokens"])
    print("mean tokens per label (PASS / FAIL, ratio):")
    for c in sorted(ln):
        mp, mf = (sum(ln[c][k]) / len(ln[c][k]) if ln[c][k] else 0 for k in ("PASS", "FAIL"))
        flag = "  <-- length gap over 15%" if mp and mf and abs(mp - mf) / max(mp, mf) > 0.15 else ""
        print(f"  {c:18s} {mp:6.1f} / {mf:6.1f}  {mp / mf if mf else 0:5.2f}{flag}")
    for d in dropped[:25]:
        print("  DROPPED", d["id"], "-", d["reason"])
    if len(dropped) > 25:
        print(f"  ... and {len(dropped) - 25} more")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fragments", required=True, help='glob, e.g. "data/fragments/claude/pilot_*.txt"')
    ap.add_argument("--out", required=True)
    ap.add_argument("--dropped", default=None)
    ap.add_argument("--sheet", default=None, help="write the human skim sheet (xlsx)")
    ap.add_argument("--source", default="claude")
    args = ap.parse_args()

    builder = Builder(args.source)
    files = sorted(glob.glob(os.path.join(REPO, args.fragments)) if not os.path.isabs(args.fragments) else glob.glob(args.fragments))
    frags = [f for p in files for f in parse_file(p)]
    dup = [i for i, c in collections.Counter(f["id"] for f in frags).items() if c > 1]
    if dup:
        sys.exit(f"duplicate fragment ids: {dup[:5]}")

    rows, dropped = [], []
    for f in frags:
        row, drop = builder.row(f)
        (rows if row else dropped).append(row or drop)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8", newline="\n") as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    if args.dropped:
        json.dump(dropped, open(args.dropped, "w", encoding="utf-8"), indent=1)
    report(frags, rows, dropped)
    if args.sheet:
        write_sheet(rows, args.sheet)


def write_sheet(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Pilot"
    arial = Font(name="Arial", size=10)
    head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    side = Side(style="thin", color="BFBFBF")
    heads = ["Item", "Check", "What the model sees", "Claude's label", "Do you agree?", "Notes (your own words; required if you disagree)"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.alignment = head, PatternFill("solid", fgColor="1F3864"), Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, 2):
        vals = [r["meta"]["id"], r["meta"]["check"], r["state"], r["meta"]["label"], None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name="Arial", size=10, bold=(j in (1, 4)))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(left=side, right=side, top=side, bottom=side)
            if j in (5, 6):
                c.fill = PatternFill("solid", fgColor="FFFF99")
    for col, w in zip("ABCDEF", (14, 16, 100, 12, 14, 44)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"
    dv = DataValidation(type="list", formula1='"AGREE,DISAGREE,UNSURE"', allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(rows) + 1}")
    wb.save(path)
    print(f"wrote skim sheet {path} ({len(rows)} items)")


if __name__ == "__main__":
    main()
