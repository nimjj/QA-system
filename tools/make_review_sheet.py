"""Build the Phase 0 review workbook: a blind-first check of the 4 Main calls' expected scorecards.

The expected grades in qa-analysis/inputs/Main tests/expected/ were written by Claude, not by a human. This
workbook lets a human grade the same 44 items (4 calls x 11 criteria) FIRST, without seeing Claude's grade,
and then compare. Excerpts come from neutral extractors over the transcript, not from Claude's cited evidence.

Sheets:
    Read me first           what to do, progress counters, one example row
    Review                  44 rows: standard, excerpt from the call, YOUR verdict (dropdown), notes
    Call 1 .. Call 4        the full transcripts (numbered lines, time, speaker, sentiment)
    Open after finishing    Claude's grades next to yours, agreement counts, and both scorecard scores

Usage (any Python with openpyxl):
    python tools/make_review_sheet.py
    python tools/make_review_sheet.py --qa-root ../qa-analysis --out data/review/main_scorecard_review.xlsx
"""
import argparse
import difflib
import glob
import json
import math
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)

FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SMALL = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFFF99")       # yellow = you fill this in
FILL_NOTE = PatternFill("solid", fgColor="F2F2F2")
FILL_WARN = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# (criterion name exactly as in criteria_config.json, short id, plain-English standard)
CRITERIA = [
    ("Branding and Survey Check", "BRAND",
     "PASS: the agent says 'Thank you for calling S-Net...' in their first 4 lines AND 'Thank you for choosing S-Net' "
     "in their last 4 lines.\nFAIL: either script is missing."),
    ("Hold time and Dead Air", "HOLD",
     "FAIL: 2 or more silences over 30 s that are NOT a managed hold (announced hold such as 'let me place you on a "
     "brief hold', or ended with 'thanks for holding').\nPASS: otherwise."),
    ("Personalized the call/ticket appropriately", "NAME",
     "PASS: the agent says the customer's name at least once.\nFAIL: never."),
    ("Empathy & Acknowledgment Statement", "EMP",
     "PASS: when the customer is upset or frustrated, the agent acknowledges the feeling (warm, specific), or the "
     "customer never gets upset.\nFAIL: any clear upset moment is ignored, brushed off or answered coldly."),
    ("Build rapport and observed professionalism", "RAPPORT",
     "PASS unless you can quote a single agent line that is rude, condescending, dismissive, sarcastic or "
     "unprofessional (no insult needed: 'that's basic stuff' counts).\nFAIL: one such line is enough."),
    ("Paraphrasing", "PARA",
     "PASS: early in the call the agent restates the customer's problem in their own words (a summary, usually "
     "checked with 'is that right?').\nFAIL: the agent only asks questions that reuse the customer's words, or "
     "never restates it."),
    ("Verified customer", "VERIFY",
     "PASS: the agent asks for the PIN, the address or a security question within the first 4 minutes.\n"
     "FAIL: not asked in that window."),
    ("Probing", "PROBE",
     "PASS: the agent asks real diagnostic questions about the problem (when it started, which device, what "
     "lights or errors, what was tried).\nFAIL: only greetings, identity checks, confirmations, upsell or closing "
     "questions."),
    ("Took ownership of the problem", "OWN",
     "PASS unless you can quote an agent line that blames another team, deflects, sends the customer elsewhere or "
     "refuses to help.\nFAIL: one such line is enough, even if the agent also helps elsewhere."),
    ("Active listening", "AL",
     "FAIL: the agent asks again for information the customer already gave (a careless repeat).\nPASS: otherwise. "
     "Re-checking something that changes ('what is the light doing now?') is NOT a repeat."),
    ("Confirmed the issue is resolved", "RESOLVED",
     "PASS: before closing, the agent explicitly confirms the issue is fixed (usually asks the customer and states "
     "it is resolved).\nFAIL: the call just ends, the agent assumes it is fixed, the customer says it still fails, "
     "or it is handed off."),
]

_FRUSTRATION_RE = re.compile(
    r"\b(frustrat\w*|annoy\w*|ridiculous|nightmare|unbelievable|fed up|sick of|embarrass\w*|stress\w*|worried"
    r"|third time|again and again|furious|angry|useless|unacceptable|cancel\w*)\b", re.IGNORECASE)
_ANSWERED_RE = re.compile(
    r"\b(like i said|as i said|i already|i told you|i mentioned|i just said|you already asked|i just told)\b", re.IGNORECASE)
_CHECKIN_RE = re.compile(
    r"\b(is that right|is that correct|did i get that right|does that make sense|anything else|still there"
    r"|still with me|can you hear me|are you there|how can i help|how are you)\b", re.IGNORECASE)
_SETUP_RE = re.compile(
    r"\b(thank you for calling|account number|pin|verify|verified|verification|name on the account|address"
    r"|date of birth|security question)\b", re.IGNORECASE)


def sec(stamp):
    """'[01:24]' or '[1:02:03]' -> seconds."""
    parts = [int(p) for p in re.findall(r"\d+", stamp)]
    return parts[0] * 60 + parts[1] if len(parts) == 2 else parts[0] * 3600 + parts[1] * 60 + parts[2]


def trim(text, n):
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def load_call(path, expected_path):
    d = json.load(open(path, encoding="utf-8"))
    e = json.load(open(expected_path, encoding="utf-8"))
    turns = []
    for i, t in enumerate(d["transcript"], 1):
        start = t.get("start_time_sec", 0)
        end = sec(t["end_time"]) if t.get("end_time") else start + 10
        turns.append({"n": i, "t": start, "end": end, "time": t.get("start_time", f"[{start // 60:02d}:{start % 60:02d}]"),
                      "spk": t["speaker"], "text": t["text"], "sent": t.get("sentiment_score", 0.0)})
    return {"name": d["customer_name"], "turns": turns, "expected": e, "file": os.path.basename(path)}


def line(t, n=170):
    return f"#{t['n']} {t['time']} {t['spk']}: {trim(t['text'], n)}"


def question_sentences(text):
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if "?" in s]


# ---------------------------------------------------------------- neutral excerpt extractors
def ex_branding(call, sheet):
    agent = [t for t in call["turns"] if t["spk"].lower() == "agent"]
    picks = agent[:4] + [t for t in agent[-4:] if t not in agent[:4]]
    return "Agent's first 4 and last 4 lines:\n" + "\n".join(line(t, 150) for t in picks)


def ex_hold(call, sheet):
    ts, out = call["turns"], []
    for a, b in zip(ts, ts[1:]):
        gap = b["t"] - a["end"]
        if gap > 30:
            out.append(f"Silence of {gap} s before #{b['n']}:\n   {line(a, 150)}\n   {line(b, 150)}")
    if not out:
        return "No silence over 30 s between turns (end of one turn to start of the next)."
    return f"{len(out)} silence(s) over 30 s" + (" (first 6 shown)" if len(out) > 6 else "") + ":\n" + "\n".join(out[:6])


def ex_name(call, sheet):
    name = call["name"]
    hits = [t for t in call["turns"] if t["spk"].lower() == "agent" and re.search(rf"\b{re.escape(name)}\b", t["text"], re.I)]
    if not hits:
        return f"The agent never says '{name}' (searched all {sum(t['spk'].lower() == 'agent' for t in call['turns'])} agent lines)."
    return f"Customer name: {name}. The agent says it in {len(hits)} line(s); first 3:\n" + "\n".join(line(t, 150) for t in hits[:3])


def ex_empathy(call, sheet):
    ts, cust_scores, cands = call["turns"], [], []
    for i, t in enumerate(ts):
        if t["spk"].lower() != "customer":
            continue
        drop = max((p - t["sent"] for p in cust_scores[-3:]), default=0.0)
        cust_scores.append(t["sent"])
        if t["sent"] <= -10 or drop >= 8 or (t["sent"] < 0 and _FRUSTRATION_RE.search(t["text"])):
            cands.append((t["sent"], -drop, i))
    picked = []
    for _, _, i in sorted(cands):                       # most upset first, windows must not overlap
        if all(abs(i - j) > 3 for j in picked):
            picked.append(i)
        if len(picked) == 3:
            break
    if not picked:
        return ("No clearly upset customer moment found (no turn with sentiment <= -10 or frustration words).\n"
                "Skim the customer's lines on the Call sheet. If they never get upset, the verdict is PASS.")
    blocks = []
    for i in sorted(picked):
        blocks.append("\n".join(line(t, 170) for t in ts[i:i + 3]) + f"\n   (customer sentiment {ts[i]['sent']:+.0f})")
    return f"{len(cands)} upset-looking customer turn(s); the {len(picked)} lowest shown, each with the next 2 turns:\n" + "\n- - -\n".join(blocks)


def ex_read_all_agent(call, sheet):
    n = sum(t["spk"].lower() == "agent" for t in call["turns"])
    return (f"Read ALL {n} agent lines. On sheet '{sheet}', filter the Speaker column to 'Agent' "
            f"(drop-down on the header). No excerpt is shown on purpose: a rude or deflecting line can hide anywhere.")


def ex_para(call, sheet):
    ts = call["turns"]
    cust = [t for t in ts if t["spk"].lower() == "customer" and len(t["text"].split()) >= 9 and not re.search(r"\d{3,}", t["text"])][:2]
    agent = [t for t in ts if t["spk"].lower() == "agent" and not _SETUP_RE.search(t["text"])][:5]
    return ("Customer's opening (account-detail turns skipped):\n" + "\n".join(line(t, 200) for t in cust)
            + "\nAgent's first 5 lines after greeting and verification:\n" + "\n".join(line(t, 190) for t in agent))


def ex_verify(call, sheet):
    early = [t for t in call["turns"] if t["spk"].lower() == "agent" and t["t"] <= 240]
    return f"Agent lines in the first 4 minutes ({len(early)}; first 8 shown):\n" + "\n".join(line(t, 150) for t in early[:8])


def ex_probe(call, sheet):
    qs = [(t, q) for t in call["turns"] if t["spk"].lower() == "agent" for q in question_sentences(t["text"]) if len(q.split()) >= 4]
    shown = qs[:12]
    return (f"Agent's questions ({len(qs)} in the call; first {len(shown)} shown). Ignore greetings and verification:\n"
            + "\n".join(f"#{t['n']} {t['time']}: {trim(q, 130)}" for t, q in shown))


def ex_al(call, sheet):
    ts = call["turns"]
    qs = [(t, q) for t in ts if t["spk"].lower() == "agent" for q in question_sentences(t["text"])
          if len(q.split()) >= 4 and not _CHECKIN_RE.search(q)]
    pairs = []
    for i, (t1, q1) in enumerate(qs):
        for t2, q2 in qs[i + 1:]:
            if t2["n"] != t1["n"] and difflib.SequenceMatcher(None, q1.lower(), q2.lower()).ratio() >= 0.55:
                pairs.append((difflib.SequenceMatcher(None, q1.lower(), q2.lower()).ratio(), t1, q1, t2, q2))
    pairs.sort(key=lambda p: -p[0])
    said = [t for t in ts if t["spk"].lower() == "customer" and _ANSWERED_RE.search(t["text"])]
    parts = []
    if pairs:
        parts.append(f"Agent questions that look alike ({len(pairs)} pair(s); top {min(4, len(pairs))}):\n" + "\n".join(
            f"#{a['n']} {a['time']}: {trim(qa, 125)}\n#{b['n']} {b['time']}: {trim(qb, 125)}" for _, a, qa, b, qb in pairs[:4]))
    if said:
        parts.append("Customer replies pointing out they already answered:\n" + "\n".join(line(t, 150) for t in said[:3]))
    if not parts:
        return ("No look-alike agent questions and no 'like I said' replies found automatically.\n"
                "Skim the agent's questions on the Call sheet for anything asked twice.")
    return "\n".join(parts)


def ex_resolved(call, sheet):
    return "Last 12 lines of the call:\n" + "\n".join(line(t, 170) for t in call["turns"][-12:])


EXTRACTORS = {"BRAND": ex_branding, "HOLD": ex_hold, "NAME": ex_name, "EMP": ex_empathy, "RAPPORT": ex_read_all_agent,
              "PARA": ex_para, "VERIFY": ex_verify, "PROBE": ex_probe, "OWN": ex_read_all_agent, "AL": ex_al,
              "RESOLVED": ex_resolved}


# ---------------------------------------------------------------- workbook helpers
def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font, cell.fill, cell.border = F_HEAD, FILL_HEAD, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def put(ws, ref, value, font=F_BODY, fill=None, wrap=True, border=True):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = WRAP if wrap else Alignment(vertical="top")
    if border:
        cell.border = BORDER
    return cell


def build(qa_root, out, subset=None):
    """Write the workbook. With `subset` (a set of item ids like '4-AL') it is the BLIND version: only those
    rows, and no sheet containing Claude's grades or evidence, so nothing can leak into the reviewer's verdicts."""
    blind = subset is not None
    main = os.path.join(qa_root, "inputs", "Main tests")
    config = json.load(open(os.path.join(qa_root, "resources", "criteria_config.json"), encoding="utf-8"))["tenant-abc"]
    meta = {}
    for cat in config["categories"]:
        for it in cat["line_items"]:
            meta[it["name"]] = (cat["name"], it["deduction_value"])
    weights = config["category_weights"]

    calls = []
    for path in sorted(glob.glob(os.path.join(main, "gold_*.json"))):
        calls.append(load_call(path, os.path.join(main, "expected", "expected_" + os.path.basename(path))))
    assert len(calls) == 4, f"expected 4 Main calls, found {len(calls)}"
    for k, c in enumerate(calls, 1):
        c["label"] = f"Call {k} - {c['name']}"
        c["sheet"] = f"Call {k}"
        c["tier"] = c["expected"]["grader"]["tier"]

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Read me first"
    ws_rev = wb.create_sheet("Review")
    ws_calls = [wb.create_sheet(c["sheet"]) for c in calls]
    ws_cmp = None if blind else wb.create_sheet("Open after finishing")

    # ---------------- Review
    heads = ["Item", "Call", "Criterion", "The standard", "Excerpt from the call (the whole call is on the Call sheets)",
             "YOUR VERDICT", "Notes (required if UNSURE or if you are unsure why)"]
    for j, h in enumerate(heads, 1):
        ws_rev.cell(1, j, h)
    style_header(ws_rev, 1, len(heads))
    rows = []   # (call, criterion, id, category, deduction)
    r = 2
    for c in calls:
        for name, short, standard in CRITERIA:
            cat, ded = meta[name]
            item_id = f"{c['sheet'][-1]}-{short}"
            if blind and item_id not in subset:
                continue
            put(ws_rev, f"A{r}", item_id, F_BOLD)
            put(ws_rev, f"B{r}", c["label"])
            put(ws_rev, f"C{r}", name, F_BOLD)
            put(ws_rev, f"D{r}", standard)
            put(ws_rev, f"E{r}", EXTRACTORS[short](c, c["sheet"]))
            put(ws_rev, f"F{r}", None, F_BOLD, FILL_INPUT)
            put(ws_rev, f"G{r}", None, F_BODY, FILL_INPUT)
            ws_rev[f"F{r}"].alignment = Alignment(horizontal="center", vertical="top")
            rows.append((c, name, item_id, cat, ded, r))
            r += 1
    last_rev = r - 1
    for col, w in zip("ABCDEFG", (9, 14, 22, 46, 100, 14, 38)):
        ws_rev.column_dimensions[col].width = w
    ws_rev.freeze_panes = "D2"
    ws_rev.auto_filter.ref = f"A1:G{last_rev}"
    dv = DataValidation(type="list", formula1='"PASS,FAIL,UNSURE"', allow_blank=True, showErrorMessage=True,
                        errorTitle="PASS, FAIL or UNSURE", error="Pick PASS, FAIL or UNSURE from the list.")
    ws_rev.add_data_validation(dv)
    dv.add(f"F2:F{last_rev}")

    # ---------------- call sheets
    for c, ws in zip(calls, ws_calls):
        for j, h in enumerate(["#", "Time", "Speaker", "Sentiment", "Text"], 1):
            ws.cell(1, j, h)
        style_header(ws, 1, 5)
        for i, t in enumerate(c["turns"], 2):
            vals = (t["n"], t["time"], t["spk"], t["sent"], t["text"])
            for j, v in enumerate(vals, 1):
                cell = ws.cell(i, j, v)
                cell.font = F_BOLD if (j == 3 and t["spk"].lower() == "agent") else F_BODY
                cell.alignment = WRAP
                cell.border = BORDER
        for col, w in zip("ABCDE", (6, 9, 11, 10, 120)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:E{len(c['turns']) + 1}"

    # ---------------- comparison sheet (Claude's grades): full version only
    if not blind:
        build_comparison(ws_cmp, calls, rows, weights)

    ws = ws0
    write_readme(ws, rows, last_rev, blind)

    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    return len(rows), last_rev


def build_comparison(cmp_, calls, rows, weights):
    cmp_.sheet_properties.tabColor = "A6A6A6"
    cmp_["A1"] = "STOP: finish the Review sheet first. This sheet shows Claude's grades; seeing them early would bias yours."
    cmp_["A1"].font, cmp_["A1"].fill = Font(name=FONT, size=11, bold=True, color="C00000"), FILL_WARN
    cmp_.merge_cells("A1:I1")
    cmp_["A2"] = ("Scores below use the real scoring rules: each category starts at 100 and loses the item's deduction "
                  "per FAIL; Soft Skills x0.333 + Technical Knowledge x0.667. Auto-fail (agent profanity) is not modelled. "
                  "A blank or UNSURE verdict counts as PASS here; we settle those together.")
    cmp_["A2"].font, cmp_["A2"].alignment = F_SMALL, WRAP
    cmp_.merge_cells("A2:I2")
    cmp_.row_dimensions[2].height = 38

    # weights block (inputs; source: criteria_config.json)
    put(cmp_, "A4", "Category weight", F_BOLD, FILL_NOTE)
    put(cmp_, "B4", "Weight", F_BOLD, FILL_NOTE)
    put(cmp_, "C4", "Source", F_BOLD, FILL_NOTE)
    put(cmp_, "A5", "Soft Skills")
    put(cmp_, "B5", weights["Soft Skills"], Font(name=FONT, size=10, color="0000FF"))
    put(cmp_, "C5", "qa-analysis/resources/criteria_config.json (category_weights)")
    put(cmp_, "A6", "Technical Knowledge")
    put(cmp_, "B6", weights["Technical Knowledge"], Font(name=FONT, size=10, color="0000FF"))
    put(cmp_, "C6", "qa-analysis/resources/criteria_config.json (category_weights)")
    cmp_.merge_cells("C4:E4"); cmp_.merge_cells("C5:E5"); cmp_.merge_cells("C6:E6")

    sum_head = 8
    for j, h in enumerate(["Call", "Tier (Claude's label)", "Score from Claude's grades", "Score from YOUR grades",
                           "Difference", "You agree", "You DISAGREE", "You are unsure", "Not reviewed yet"], 1):
        cmp_.cell(sum_head, j, h)
    style_header(cmp_, sum_head, 9)
    det_head = sum_head + len(calls) + 4
    d0 = det_head + 1
    d1 = d0 + len(rows) - 1
    rng = lambda col: f"${col}${d0}:${col}${d1}"

    def score(verdict_col, row):
        parts = []
        for cat, wcell in (("Soft Skills", "$B$5"), ("Technical Knowledge", "$B$6")):
            parts.append(f'MAX(0,100-SUMIFS({rng("E")},{rng("B")},$A{row},{rng("D")},"{cat}",{rng(verdict_col)},"FAIL"))*{wcell}')
        return "=ROUND(" + "+".join(parts) + ",1)"

    for k, c in enumerate(calls):
        rr = sum_head + 1 + k
        put(cmp_, f"A{rr}", c["label"], F_BOLD)
        put(cmp_, f"B{rr}", c["tier"])
        put(cmp_, f"C{rr}", score("F", rr))
        put(cmp_, f"D{rr}", score("G", rr))
        put(cmp_, f"E{rr}", f"=D{rr}-C{rr}")
        put(cmp_, f"F{rr}", f'=COUNTIFS({rng("B")},$A{rr},{rng("H")},"agree")')
        put(cmp_, f"G{rr}", f'=COUNTIFS({rng("B")},$A{rr},{rng("H")},"DISAGREE")')
        put(cmp_, f"H{rr}", f'=COUNTIFS({rng("B")},$A{rr},{rng("H")},"unsure")')
        put(cmp_, f"I{rr}", f"={len(CRITERIA)}-F{rr}-G{rr}-H{rr}")
    tot = sum_head + 1 + len(calls)
    put(cmp_, f"A{tot}", "All 4 calls", F_BOLD, FILL_NOTE)
    for col in "BCDE":
        put(cmp_, f"{col}{tot}", None, F_BODY, FILL_NOTE)
    for col in "FGHI":
        put(cmp_, f"{col}{tot}", f"=SUM({col}{sum_head + 1}:{col}{tot - 1})", F_BOLD, FILL_NOTE)

    for j, h in enumerate(["Item", "Call", "Criterion", "Category", "Deduction if FAIL", "Claude's grade", "Your verdict",
                           "Result", "Claude's evidence (as written when grading)"], 1):
        cmp_.cell(det_head, j, h)
    style_header(cmp_, det_head, 9)
    for n, (c, name, item_id, cat, ded, rev_row) in enumerate(rows):
        rr = d0 + n
        item = next(i for i in c["expected"]["grader"]["items"] if i["name"] == name)
        put(cmp_, f"A{rr}", item_id, F_BOLD)
        put(cmp_, f"B{rr}", c["label"])
        put(cmp_, f"C{rr}", name)
        put(cmp_, f"D{rr}", cat)
        put(cmp_, f"E{rr}", ded)
        put(cmp_, f"F{rr}", item["rating"], F_BOLD)
        put(cmp_, f"G{rr}", f'=IF(Review!F{rev_row}="","",Review!F{rev_row})', F_BOLD)
        put(cmp_, f"H{rr}", f'=IF(G{rr}="","",IF(G{rr}="UNSURE","unsure",IF(G{rr}=F{rr},"agree","DISAGREE")))')
        put(cmp_, f"I{rr}", item["evidence"])
    for col, w in zip("ABCDEFGHI", (24, 22, 28, 20, 12, 14, 14, 12, 90)):
        cmp_.column_dimensions[col].width = w


def write_readme(ws, rows, last_rev, blind):
    n = len(rows)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 44
    ws["A1"] = "Phase 0: grade these call items yourself"
    ws["A1"].font = F_TITLE
    common = [
        ("How to decide", "Each row has the standard (PASS if / FAIL if) and an excerpt from the call. The excerpt is a "
                          "starting point, not the whole evidence: lines are numbered (#12) so you can find them on the "
                          "'Call 1' ... 'Call 4' sheets. For Rapport and Ownership you must read all agent lines."),
        ("UNSURE", "Use it when a reasonable person could go either way. Write one line in Notes saying why, in your own "
                   "words. We will settle those together."),
    ]
    if blind:
        text = [
            ("Why", "The expected grades for the 4 Main calls were written by Claude, not by a person, and Ollama and Laya "
                    "are both measured against them. A human verdict is needed on the items that decide the conclusions."),
            ("What you do", f"On the 'Review' sheet, give each of the {n} rows a verdict in the yellow YOUR VERDICT column "
                            "(PASS, FAIL or UNSURE). The rows are a mix of items from the 4 calls; the other items of "
                            "those calls are not in this file."),
            ("Please", "Grade these yourself and read the calls. This file is only useful as an independent check if the "
                       "verdicts and notes are your own judgement. There is no answer sheet in this file on purpose."),
        ] + common + [
            ("Time", "About 45 minutes. You can stop and resume any time; progress is saved with the file."),
            ("When done", "Save the file and tell Claude. Claude compares your verdicts with the existing grades and "
                          "recomputes Ollama's and Laya's error."),
        ]
    else:
        text = [
            ("Why", "The expected grades for these 4 calls were written by Claude, not by a person. Everything we measure "
                    "(is Laya better than Ollama?) is judged against them, so they need a human's verdict first."),
            ("What you do", f"On the 'Review' sheet, give each of the {n} rows a verdict in the yellow YOUR VERDICT column "
                            "(PASS, FAIL or UNSURE). 4 calls x 11 criteria."),
            ("Grade first, compare later", f"Do NOT open 'Open after finishing' until all {n} are done. It shows Claude's "
                                           "grades, and seeing them first would nudge yours."),
        ] + common + [
            ("Time", "About 1.5 to 2 hours. You can stop and resume any time; progress is saved with the file."),
            ("When done", "Save the file and tell Claude. Claude reads it, compares, and recomputes Ollama's and Laya's error "
                          "against YOUR grades."),
        ]
    r0 = 3
    for i, (k, v) in enumerate(text):
        put(ws, f"A{r0 + i}", k, F_BOLD, FILL_NOTE)
        put(ws, f"B{r0 + i}", v)
        ws.merge_cells(f"B{r0 + i}:E{r0 + i}")
        ws.row_dimensions[r0 + i].height = 15 * max(2, math.ceil(len(v) / 78))
    rp = r0 + len(text) + 1
    put(ws, f"A{rp}", "Progress", F_BOLD, FILL_NOTE)
    put(ws, f"B{rp}", "Verdicts given", F_BOLD, FILL_NOTE)
    put(ws, f"C{rp}", f"=COUNTA(Review!F2:F{last_rev})")
    put(ws, f"D{rp}", f"of {len(rows)}")
    put(ws, f"B{rp + 1}", "Marked UNSURE", F_BOLD, FILL_NOTE)
    put(ws, f"C{rp + 1}", f'=COUNTIF(Review!F2:F{last_rev},"UNSURE")')
    put(ws, f"B{rp + 2}", "Still to do", F_BOLD, FILL_NOTE)
    put(ws, f"C{rp + 2}", f"={len(rows)}-C{rp}")
    for rr in range(rp, rp + 3):
        ws.cell(rr, 1).border = BORDER
    ws.merge_cells(f"A{rp}:A{rp + 2}")

    rl = rp + 4
    put(ws, f"A{rl}", "Legend", F_BOLD, FILL_NOTE)
    put(ws, f"B{rl}", "Yellow cells", F_BOLD, FILL_INPUT)
    put(ws, f"C{rl}", "are the ones you fill in (verdict and notes). Nothing else needs editing.")
    ws.merge_cells(f"C{rl}:E{rl}")
    put(ws, f"A{rl + 1}", "Example of a filled row", F_BOLD, FILL_NOTE)
    ex = [("Item", "9-OWN"), ("Criterion", "Took ownership of the problem"), ("YOUR VERDICT", "FAIL"),
          ("Notes", "#58 'That's the billing team's problem, call them' deflects. Rest of the call was fine.")]
    for i, (k, v) in enumerate(ex):
        put(ws, f"B{rl + 1 + i}", k, F_BOLD, FILL_NOTE)
        put(ws, f"C{rl + 1 + i}", v, F_BODY)
        ws.merge_cells(f"C{rl + 1 + i}:E{rl + 1 + i}")
    put(ws, f"A{rl + 5}", "(illustration only: not a real item)", F_SMALL, border=False)
    ws.merge_cells(f"A{rl + 1}:A{rl + 4}")
    ws.cell(rl + 1, 1).border = BORDER


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--qa-root", default=os.path.join(REPO, "..", "qa-analysis"))
    ap.add_argument("--out", default=None)
    ap.add_argument("--subset", default=None,
                    help="subset_main.json from pick_review_subset.py: build the BLIND subset workbook (no answers)")
    args = ap.parse_args()
    subset = set(json.load(open(args.subset, encoding="utf-8"))["items"]) if args.subset else None
    out = args.out or os.path.join(REPO, "data", "review",
                                   "main_review_subset_blind.xlsx" if subset else "main_scorecard_review.xlsx")
    n, last = build(os.path.abspath(args.qa_root), os.path.abspath(out), subset)
    print(f"wrote {out}: {n} review rows (Review!A2:G{last})" + (" [blind: no answer sheet]" if subset else ""))


if __name__ == "__main__":
    main()
